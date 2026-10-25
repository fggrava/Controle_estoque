import io, uuid
from decimal import Decimal
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from werkzeug.security import generate_password_hash
import app as module
from models import Base, User, State, Material, Attendance, Item, Audit
from services import excel, HEADERS, parse_base, same_balances


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.delenv("VERCEL", raising=False)
    monkeypatch.setenv("APP_ORIGIN", "http://localhost")
    import os

    engine = create_engine(
        os.environ.get("TEST_DATABASE_URL") or "sqlite:///" + str(tmp_path / "test.db")
    )
    if os.environ.get("TEST_DATABASE_URL"):
        Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    module.engine = engine
    with Session(engine) as s:
        s.add(State(id=1))
        s.add_all(
            [
                User(
                    login="admin",
                    name="Administrador",
                    password=generate_password_hash("SenhaAdmin123!"),
                    role="admin",
                ),
                User(
                    login="operador",
                    name="Operador",
                    password=generate_password_hash("SenhaOperador123!"),
                ),
            ]
        )
        s.commit()
    with module.app.test_client() as c:
        yield c
    engine.dispose()
    module.engine = None


def login(c, name="admin"):
    r = c.post(
        "/api/login",
        json={
            "login": name,
            "password": "SenhaAdmin123!" if name == "admin" else "SenhaOperador123!",
        },
        headers={"Origin": "http://localhost"},
    )
    assert r.status_code == 200, r.json
    return {"Origin": "http://localhost", "X-CSRF-Token": r.json["csrf"]}


def base(q=10, code="000123", depot="0101"):
    return excel(
        {
            "Base": [
                HEADERS,
                [
                    code,
                    "Peça teste",
                    "UN",
                    "1002",
                    depot,
                    "Consumo",
                    "",
                    "VBRA",
                    q,
                    q * 5,
                ],
            ]
        }
    ).read()


def imp(c, h, raw, confirm=False, extra=None):
    if not confirm:
        return c.post(
            "/api/import/preview",
            data={"file": (io.BytesIO(raw), "base.xlsx")},
            headers=h,
        )
    p = imp(c, h, raw).json
    form = {k: str(p[k]) for k in ["cycle", "revision", "digest"]}
    form.update(reconciled="true", confirm_same="true")
    form.update(extra or {})
    form["file"] = (io.BytesIO(raw), "base.xlsx")
    return c.post("/api/import/confirm", data=form, headers=h)


def launch(c, h, q="2", cid=None, cycle=1):
    mid = c.get("/api/materials").json[0]["id"]
    return c.post(
        "/api/attendances",
        json={
            "client_id": cid or str(uuid.uuid4()),
            "cycle": cycle,
            "os": "700691519",
            "receiver": "Douglas",
            "confirm_duplicate": True,
            "items": [{"material_id": mid, "quantity": q}],
        },
        headers=h,
    )


def test_auth_and_csrf(client):
    assert client.get("/api/materials").status_code == 401
    h = login(client)
    assert (
        client.post(
            "/api/users", json={}, headers={"Origin": "https://evil.test"}
        ).status_code
        == 403
    )
    assert (
        client.post(
            "/api/users", json={}, headers={"Origin": "http://localhost"}
        ).status_code
        == 403
    )
    h = login(client, "operador")
    assert client.get("/api/users").status_code == 403
    assert client.post("/api/import/preview", headers=h).status_code == 403


def test_import_export_correction_cycle_and_equal_saldo(client):
    h = login(client)
    raw = base()
    assert imp(client, h, raw, True).status_code == 200
    a = launch(client, h)
    assert a.status_code == 200, a.json
    assert client.get("/api/materials").json[0]["balance"] == "8.000"
    assert imp(client, h, base(9), True).status_code == 409
    assert (
        client.post(
            "/api/export",
            json={"full": False, "start": "2020-01-01", "end": "2030-01-01"},
            headers=h,
        ).status_code
        == 200
    )
    assert imp(client, h, base(9), True).status_code == 409
    assert client.post("/api/export", json={"full": True}, headers=h).status_code == 200
    mid = client.get("/api/materials").json[0]["id"]
    edit = client.patch(
        "/api/attendances/" + str(a.json["id"]),
        json={
            "revision": 1,
            "reason": "Correção",
            "items": [{"material_id": mid, "quantity": "3"}],
        },
        headers=h,
    )
    assert edit.status_code == 200, edit.json
    assert imp(client, h, base(9), True).status_code == 409
    client.post("/api/export", json={"full": True}, headers=h)
    assert imp(client, h, raw).json["same"] is True
    assert imp(client, h, raw, True, {"confirm_same": "false"}).status_code == 409
    assert imp(client, h, base(9), True).status_code == 200
    assert client.get("/api/materials").json[0]["balance"] == "9.000"
    assert launch(client, h, cycle=1).status_code == 409


def test_idempotency_and_rollback(client):
    h = login(client)
    imp(client, h, base(), True)
    cid = str(uuid.uuid4())
    assert launch(client, h, cid=cid).status_code == 200
    assert launch(client, h, cid=cid).json["duplicate"]
    assert len(client.get("/api/attendances").json) == 1
    assert launch(client, h, "99").status_code == 409
    assert client.get("/api/materials").json[0]["balance"] == "8.000"
    mid = client.get("/api/materials").json[0]["id"]
    r = client.post(
        "/api/attendances",
        json={
            "client_id": str(uuid.uuid4()),
            "cycle": 1,
            "os": "123",
            "receiver": "X",
            "items": [
                {"material_id": mid, "quantity": "2"},
                {"material_id": mid, "quantity": "1"},
            ],
        },
        headers=h,
    )
    assert r.status_code == 400
    assert client.get("/api/materials").json[0]["balance"] == "8.000"


def test_ownership_and_filters(client):
    h = login(client)
    imp(client, h, base(), True)
    aid = launch(client, h).json["id"]
    h = login(client, "operador")
    assert client.get("/api/attendances").json == []
    assert (
        client.patch(
            f"/api/attendances/{aid}",
            json={"revision": 1, "cancel": True, "reason": "teste"},
            headers=h,
        ).status_code
        == 403
    )
    launch(client, h)
    from datetime import datetime
    from zoneinfo import ZoneInfo

    day = datetime.now(ZoneInfo("America/Sao_Paulo")).date().isoformat()
    d = client.get(f"/api/dashboard?start={day}&end={day}&user=1").json
    assert d["attendances"] == 1
    h = login(client)
    d = client.get(f"/api/dashboard?start={day}&end={day}&user=1").json
    assert d["attendances"] == 1
    assert (
        client.get("/api/dashboard?start=2020-01-02&end=2020-01-01").status_code == 400
    )


def test_parse_and_formula_protection():
    rows, _ = parse_base(base())
    assert rows[0]["code"] == "000123"
    assert rows[0]["depot"] == "0101"
    assert same_balances({"x": "1.000"}, {"x": "1"})
    assert not same_balances({"x": "1"}, {"y": "1"})
    from openpyxl import load_workbook

    assert (
        load_workbook(excel({"Teste": [['=HYPERLINK("evil")']]})).active["A1"].data_type
        == "s"
    )


def test_template_and_security_headers(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "script-src 'self'" in r.headers["Content-Security-Policy"]
    h = login(client)
    r = client.get("/api/import/template")
    assert r.status_code == 200
    assert client.get("/sw.js").headers["Service-Worker-Allowed"] == "/"
    assert client.get("/api/state").headers["Cache-Control"] == "no-store"


def test_stale_edit_and_inventory(client):
    h = login(client)
    imp(client, h, base(), True)
    aid = launch(client, h).json["id"]
    mid = client.get("/api/materials").json[0]["id"]
    assert (
        client.patch(
            f"/api/attendances/{aid}",
            json={"revision": 0, "reason": "teste", "cancel": True},
            headers=h,
        ).status_code
        == 409
    )
    assert (
        client.post(
            f"/api/inventory/{mid}",
            json={"cycle": 1, "expected": "10", "counted": "9", "reason": "contagem"},
            headers=h,
        ).status_code
        == 409
    )
    assert (
        client.post(
            f"/api/inventory/{mid}",
            json={"cycle": 1, "expected": "8", "counted": "9", "reason": "contagem"},
            headers=h,
        ).status_code
        == 200
    )
    assert imp(client, h, base(), True).status_code == 409


def test_login_throttle(client):
    for _ in range(10):
        assert (
            client.post(
                "/api/login",
                json={"login": "inexistente", "password": "errada"},
                headers={"Origin": "http://localhost"},
            ).status_code
            == 401
        )
    assert (
        client.post(
            "/api/login",
            json={"login": "inexistente", "password": "errada"},
            headers={"Origin": "http://localhost"},
        ).status_code
        == 429
    )


def test_duplicate_warning_and_historical_description(client):
    h = login(client)
    imp(client, h, base(), True)
    launch(client, h)
    mid = client.get("/api/materials").json[0]["id"]
    r = client.post(
        "/api/attendances",
        json={
            "client_id": str(uuid.uuid4()),
            "cycle": 1,
            "os": "700691519",
            "receiver": "Outro",
            "items": [{"material_id": mid, "quantity": "1"}],
        },
        headers=h,
    )
    assert r.status_code == 409 and "duplicidade" in r.json["error"]
    assert (
        client.patch(
            f"/api/materials/{mid}",
            json={"description": "Descrição alterada"},
            headers=h,
        ).status_code
        == 200
    )
    assert (
        client.get("/api/attendances").json[0]["items"][0]["description"]
        == "Peça teste"
    )


def test_user_management_and_material_create(client):
    h = login(client)
    assert (
        client.post(
            "/api/users",
            json={
                "login": "novo",
                "name": "Novo",
                "password": "SenhaNova123!",
                "role": "user",
            },
            headers=h,
        ).status_code
        == 200
    )
    uid = next(u["id"] for u in client.get("/api/users").json if u["login"] == "novo")
    assert (
        client.patch(
            f"/api/users/{uid}",
            json={"login": "renomeado", "name": "Nome", "active": True, "role": "user"},
            headers=h,
        ).status_code
        == 200
    )
    assert client.delete(f"/api/users/{uid}", headers=h).status_code == 200
    assert (
        client.post(
            "/api/materials",
            json={
                "code": "0001",
                "center": "1002",
                "depot": "0101",
                "description": "Material",
                "unit": "UN",
                "balance": "2",
                "unit_cost": "1",
            },
            headers=h,
        ).status_code
        == 200
    )
    mid = client.get("/api/materials").json[0]["id"]
    assert client.delete(f"/api/materials/{mid}", headers=h).status_code == 409
    assert imp(client, h, base(), True).status_code == 409


def test_postgres_concurrent_withdrawals(client):
    import os

    if not os.environ.get("TEST_DATABASE_URL"):
        pytest.skip("Requer PostgreSQL isolado; executado pelo GitHub Actions.")
    from concurrent.futures import ThreadPoolExecutor

    h = login(client)
    imp(client, h, base(3), True)

    def worker(_):
        with module.app.test_client() as c:
            headers = login(c)
            return launch(c, headers, "2").status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(worker, range(2)))
    assert sorted(results) == [200, 409]
    assert client.get("/api/materials").json[0]["balance"] == "1.000"


def test_production_tls_uses_bundled_ca(monkeypatch):
    import certifi
    from unittest.mock import Mock

    captured = {}
    monkeypatch.setenv("VERCEL", "1")
    monkeypatch.setenv(
        "DATABASE_URL", "postgresql://example:placeholder@db.example.test/test"
    )
    monkeypatch.setattr(module, "engine", None)

    def factory(url, **kwargs):
        captured.update(kwargs)
        return Mock()

    monkeypatch.setattr(module, "create_engine", factory)
    monkeypatch.setattr(module, "Session", lambda engine: Mock())
    with module.app.app_context():
        module.db()
    assert captured["connect_args"]["sslmode"] == "verify-full"
    assert captured["connect_args"]["sslrootcert"] == certifi.where()


def test_mb52_header_after_title_on_second_sheet():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Capa"
    wb.active.append(["Relatório SAP"])
    ws = wb.create_sheet("MB52")
    ws.append(["Posição de estoque"])
    ws.append([h.replace(" ", "\n", 1) for h in HEADERS])
    ws.append(["000123", "Peça", "UN", "1002", "0101", "Consumo", "", "VBRA", 10, 50])
    buf = io.BytesIO()
    wb.save(buf)
    rows, _ = parse_base(buf.getvalue())
    assert rows[0]["base"] == 10
    assert rows[0]["stock_value"] == 50
