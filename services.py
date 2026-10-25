import json

"""Validações de negócio independentes da interface."""

import hashlib, io, unicodedata, zipfile
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook, Workbook
from werkzeug.exceptions import BadRequest

HEADERS = [
    "Material",
    "Texto breve material",
    "Unid.medida básica",
    "Centro",
    "Depósito",
    "Denominação depósito",
    "Marc.p/elim.nív.dep.",
    "Tipo de material",
    "Utilização livre",
    "Val.utiliz.livre",
]


def normalized_text(v):
    return (
        "".join(
            c
            for c in unicodedata.normalize("NFD", str(v or ""))
            if unicodedata.category(c) != "Mn"
        )
        .lower()
        .strip()
    )


def norm(v):
    return " ".join(normalized_text(v).split())


def number(v):
    try:
        if isinstance(v, str):
            v = v.replace(".", "").replace(",", ".") if "," in v else v
        n = Decimal(str(v))
        if not n.is_finite() or abs(n) > Decimal("999999999999"):
            raise ValueError()
        return n
    except (InvalidOperation, ValueError, TypeError):
        raise BadRequest("Número inválido.")


def quantity(v):
    n = number(v)
    if n <= 0 or n != n.quantize(Decimal(".001")):
        raise BadRequest("Quantidade deve ser positiva, com até 3 casas decimais.")
    return n


def parse_base(raw):
    try:
        z = zipfile.ZipFile(io.BytesIO(raw))
        if sum(x.file_size for x in z.infolist()) > 40_000_000:
            raise BadRequest("Planilha expandida excede 40 MB.")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = None
        header_line = 0
        for ws in wb.worksheets:
            candidate = ws.iter_rows(values_only=True)
            for header_line, row in enumerate(candidate, 1):
                hdr = [norm(c) for c in row]
                if all(norm(h) in hdr for h in HEADERS):
                    rows = candidate
                    break
                if header_line >= 30:
                    break
            if rows is not None:
                break
        if rows is None:
            raise BadRequest(
                "Cabeçalhos obrigatórios nas primeiras 30 linhas: " + ", ".join(HEADERS)
            )
        ix = [hdr.index(norm(h)) for h in HEADERS]
        out = []
        seen = set()
        for line, row in enumerate(rows, header_line + 1):
            if len(out) >= 10000:
                raise BadRequest("Limite: 10.000 materiais por base.")
            if not any(v is not None for v in row):
                continue
            v = [row[i] if i < len(row) else None for i in ix]
            text = lambda i: str(v[i]).strip() if v[i] is not None else ""
            if any(
                len(text(i)) > limit
                for i, limit in [(0, 80), (1, 150), (2, 20), (3, 30), (4, 30)]
            ):
                raise BadRequest(f"Linha {line}: campo excede o tamanho permitido.")
            code = text(0)
            center = text(3)
            depot = text(4).zfill(4)
            key = (code, center, depot)
            if not code or not center or not text(1) or not text(2):
                raise BadRequest(f"Linha {line}: campos obrigatórios vazios.")
            if key in seen:
                raise BadRequest(f"Linha {line}: material/centro/depósito duplicado.")
            seen.add(key)
            q = number(v[8])
            value = number(v[9])
            if q < 0 or value < 0 or (q == 0 and value != 0):
                raise BadRequest(f"Linha {line}: saldo/valor incompatíveis.")
            if q != q.quantize(Decimal(".001")):
                raise BadRequest(f"Linha {line}: saldo com mais de 3 casas.")
            out.append(
                dict(
                    code=code,
                    description=text(1),
                    unit=text(2),
                    center=center,
                    depot=depot,
                    depot_name=text(5),
                    deletion=text(6),
                    kind=text(7),
                    base=q,
                    stock_value=value,
                    unit_cost=(value / q if q else Decimal(0)).quantize(
                        Decimal(".000001")
                    ),
                )
            )
        wb.close()
        if not out:
            raise BadRequest("Planilha vazia.")
        return out, hashlib.sha256(raw).hexdigest()
    except BadRequest:
        raise
    except Exception:
        raise BadRequest("Não foi possível ler o XLSX. Confira o arquivo.")


def snapshot(rows):
    return {
        json.dumps([r["code"], r["center"], r["depot"]], ensure_ascii=False): str(
            r["base"]
        )
        for r in rows
    }


def same_balances(a, b):
    return a.keys() == b.keys() and all(Decimal(a[k]) == Decimal(b[k]) for k in a)


def excel(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            safe = []
            for v in row:
                if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
                    v = "'" + v
                safe.append(float(v) if isinstance(v, Decimal) else v)
            ws.append(safe)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                50, max(15, max(len(str(c.value or "")) for c in col) + 2)
            )
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    return b
